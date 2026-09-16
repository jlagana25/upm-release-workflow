#!/usr/bin/env python3
"""Prepare, validate, and upload BMAT custom-content deliveries.

The remote completion marker is deliberately created only after every package
file has been uploaded and verified by size.  Accepted catalogues are tracked
in a local ledger because the Domo custom-release report has no delivery-state
column.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import os
import posixpath
import re
import shutil
import tempfile
import time
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Iterable

from auth_manager import load_bmat_sftp_credentials
from config import BMAT_BASE, DAMS_PROFILE_DIR, PRIVATE_STATE_DIR, ReleaseContext


BMAT_HOST = "sftp.content.bmat.com"
BMAT_PORT = 22
BMAT_SUBMISSION_CARD_ID = "1378540355"
BMAT_SUBMISSION_PAGE_ID = "926400857"
BMAT_RELEASES_CARD_ID = "190170615"
BMAT_LEDGER_VERSION = 1
FINAL_STATUSES = {"accepted"}
REMOTE_PENDING_STATUSES = {"ingestion_pending"}
DAMS_BASE_URL = "https://dams.universalproductionmusic.com"
DAMS_ALBUMS_URL = f"{DAMS_BASE_URL}/albums"
DAMS_AUTH_TIMEOUT_MS = 90_000
DAMS_DOWNLOAD_TIMEOUT_MS = 15 * 60 * 1000


BMAT_CARD_CONFIGS: list[dict] = [
    {
        "key": "bmat_releases",
        "card_id": BMAT_RELEASES_CARD_ID,
        "description": "Custom BMAT Releases",
        "output_fn": lambda ctx: ctx.bmat_releases_csv,
        # This link is not hosted under a fixed dashboard page.
        "page_id": None,
    },
    {
        "key": "bmat_submission",
        "card_id": BMAT_SUBMISSION_CARD_ID,
        "description": "BMAT Custom Content Production Submissions",
        "output_fn": lambda ctx: ctx.bmat_submission_xlsx,
        "page_id": BMAT_SUBMISSION_PAGE_ID,
        "format": "xlsx",
        # This is the full ingestion inventory. Selection happens by catalogue
        # after the date-filtered releases card identifies the current batch.
        "skip_timeframe": True,
    },
]


@dataclass(frozen=True)
class PreparedDelivery:
    batch_id: str
    package_dir: Path
    metadata_path: Path
    catalogues: tuple[str, ...]
    track_count: int


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _norm_header(value: object) -> str:
    return "".join(ch.lower() for ch in str(value or "") if ch.isalnum())


def _column(headers: list[str], *candidates: str) -> int:
    mapped = {_norm_header(value): index for index, value in enumerate(headers)}
    for candidate in candidates:
        if _norm_header(candidate) in mapped:
            return mapped[_norm_header(candidate)]
    raise ValueError(f"Missing required column; expected one of {candidates!r}")


def _load_xlsx_rows(path: Path) -> tuple[list[str], list[tuple[object, ...]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    # Domo sometimes writes an incorrect A1-only dimension.  Resetting it is
    # required to expose the actual 123-column custom-ingestion worksheet.
    worksheet.reset_dimensions()
    iterator = worksheet.iter_rows(values_only=True)
    try:
        headers = [str(value or "").strip() for value in next(iterator)]
    except StopIteration:
        workbook.close()
        return [], []
    rows = [
        tuple(row[: len(headers)]) + (None,) * max(0, len(headers) - len(row))
        for row in iterator
        if any(v is not None for v in row)
    ]
    workbook.close()
    return headers, rows


def _load_table(path: Path) -> tuple[list[str], list[tuple[object, ...]]]:
    if path.suffix.lower() == ".xlsx":
        return _load_xlsx_rows(path)
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        try:
            headers = [str(value or "").strip() for value in next(reader)]
        except StopIteration:
            return [], []
        return headers, [
            tuple(row[: len(headers)]) + (None,) * max(0, len(headers) - len(row))
            for row in reader
            if any(str(v).strip() for v in row)
        ]


def release_catalogues(path: Path) -> list[str]:
    headers, rows = _load_table(path)
    # A date-filtered Domo card can legitimately export no records. Depending
    # on export format that appears as a zero-row workbook, an empty/BOM-only
    # CSV, or a header-only table. All are a successful no-op for BMAT.
    if not rows:
        return []
    cat_index = _column(headers, "AlbumNo", "CatNo", "Catalog")
    seen: set[str] = set()
    result: list[str] = []
    for row in rows:
        value = str(row[cat_index] or "").strip() if cat_index < len(row) else ""
        if value and value not in seen:
            result.append(value)
            seen.add(value)
    return result


def load_ledger(path: Path) -> dict:
    if not path.exists():
        return {"version": BMAT_LEDGER_VERSION, "deliveries": []}
    data = json.loads(path.read_text(encoding="utf-8"))
    if data.get("version") != BMAT_LEDGER_VERSION or not isinstance(data.get("deliveries"), list):
        raise ValueError(f"Unsupported or malformed BMAT ledger: {path}")
    return data


def save_ledger(path: Path, ledger: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(json.dumps(ledger, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    temporary.replace(path)


def accepted_catalogues(ledger: dict) -> set[str]:
    accepted: set[str] = set()
    for delivery in ledger.get("deliveries", []):
        if delivery.get("status") in FINAL_STATUSES:
            accepted.update(str(value) for value in delivery.get("catalogues", []))
    return accepted


def pending_catalogues(releases_path: Path, ledger: dict) -> list[str]:
    unavailable = accepted_catalogues(ledger)
    for delivery in ledger.get("deliveries", []):
        if delivery.get("status") in REMOTE_PENDING_STATUSES:
            unavailable.update(str(value) for value in delivery.get("catalogues", []))
    return [
        catalogue for catalogue in release_catalogues(releases_path)
        if catalogue not in unavailable
    ]


def _prepared_from_ledger(root: Path, delivery: dict) -> PreparedDelivery:
    batch_id = str(delivery["batch_id"])
    package = root / batch_id
    metadata_matches = sorted(package.glob(f"{batch_id[:8]}_*.xlsx"))
    if len(metadata_matches) != 1:
        raise ValueError(
            f"BMAT batch {batch_id} must contain exactly one metadata workbook; "
            f"found {len(metadata_matches)}"
        )
    return PreparedDelivery(
        batch_id=batch_id,
        package_dir=package,
        metadata_path=metadata_matches[0],
        catalogues=tuple(str(value) for value in delivery.get("catalogues", [])),
        track_count=int(delivery.get("track_count", 0)),
    )


def _safe_remote_path(value: object, *, field: str, row_number: int) -> PurePosixPath | None:
    text = str(value or "").strip().replace("\\", "/")
    if not text:
        return None
    path = PurePosixPath(text)
    if path.is_absolute() or ".." in path.parts or len(path.parts) < 2:
        raise ValueError(f"Row {row_number}: unsafe {field} path {text!r}")
    return path


def _share_sum(row: tuple[object, ...], indices: list[int]) -> float | None:
    values: list[float] = []
    for index in indices:
        if index >= len(row) or row[index] in (None, ""):
            continue
        try:
            values.append(float(row[index]))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Invalid share value {row[index]!r}") from exc
    return sum(values) if values else None


def select_and_validate_submission(
    submission_path: Path,
    catalogues: Iterable[str],
) -> tuple[list[str], list[tuple[object, ...]]]:
    headers, all_rows = _load_xlsx_rows(submission_path)
    selected_catalogues = set(catalogues)
    cat_index = _column(headers, "CatNo", "AlbumNo")
    file_index = _column(headers, "TrackFilepath")
    id_index = _column(headers, "InternalID")
    artwork_index = _column(headers, "CDArtwork")
    composer_shares = [i for i, name in enumerate(headers) if name.startswith("COM:") and name.endswith(":PerformanceShare")]
    publisher_shares = [i for i, name in enumerate(headers) if name.startswith("PUB:") and name.endswith(":PerformanceShare")]

    rows: list[tuple[object, ...]] = []
    found: set[str] = set()
    internal_ids: set[str] = set()
    track_paths: set[str] = set()
    for source_row_number, row in enumerate(all_rows, start=2):
        catalogue = str(row[cat_index] or "").strip()
        if catalogue not in selected_catalogues:
            continue
        found.add(catalogue)
        internal_id = str(row[id_index] or "").strip()
        if not internal_id or internal_id in internal_ids:
            raise ValueError(f"Row {source_row_number}: InternalID is blank or duplicated")
        internal_ids.add(internal_id)

        track_path = _safe_remote_path(row[file_index], field="TrackFilepath", row_number=source_row_number)
        if track_path is None or track_path.parts[0] != catalogue:
            raise ValueError(
                f"Row {source_row_number}: TrackFilepath must begin with {catalogue}/"
            )
        if track_path.suffix.lower() != ".wav" or str(track_path) in track_paths:
            raise ValueError(f"Row {source_row_number}: invalid or duplicate WAV path")
        track_paths.add(str(track_path))

        artwork = _safe_remote_path(row[artwork_index], field="CDArtwork", row_number=source_row_number)
        if artwork is not None and artwork.parts[0] != catalogue:
            raise ValueError(f"Row {source_row_number}: CDArtwork must begin with {catalogue}/")

        for label, indices in (("composer", composer_shares), ("publisher", publisher_shares)):
            total = _share_sum(row, indices)
            if total is not None and abs(total - 100.0) > 0.001:
                raise ValueError(
                    f"Row {source_row_number}: {label} shares total {total:g}, expected 100"
                )
        rows.append(row)

    missing = selected_catalogues - found
    if missing:
        raise ValueError(f"Submission export has no rows for catalogue(s): {sorted(missing)}")
    if not rows:
        raise ValueError("No BMAT submission rows were selected")
    return headers, rows


def _next_batch_id(root: Path, ledger: dict, delivery_date: date) -> str:
    prefix = delivery_date.strftime("%Y%m%d") + "_"
    used: set[int] = set()
    for path in root.glob(prefix + "[0-9][0-9][0-9][0-9]") if root.exists() else ():
        try:
            used.add(int(path.name.rsplit("_", 1)[1]))
        except ValueError:
            pass
    for delivery in ledger.get("deliveries", []):
        batch_id = str(delivery.get("batch_id", ""))
        if batch_id.startswith(prefix):
            try:
                used.add(int(batch_id.rsplit("_", 1)[1]))
            except ValueError:
                pass
    sequence = 1
    while sequence in used:
        sequence += 1
    return f"{prefix}{sequence:04d}"


def _write_metadata(path: Path, headers: list[str], rows: list[tuple[object, ...]]) -> None:
    from openpyxl import Workbook

    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("Data")
    worksheet.append(headers)
    for row in rows:
        worksheet.append(list(row))
    temporary = path.with_name(f".{path.name}.tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(path)


def _write_dams_download_manifest(
    path: Path, headers: list[str], rows: list[tuple[object, ...]]
) -> None:
    cat_index = _column(headers, "CatNo", "AlbumNo")
    title_index = _column(headers, "CDTitle")
    track_index = _column(headers, "TrackTitle", "TrackDisplayTitle")
    number_index = _column(headers, "TrackNo")
    id_index = _column(headers, "InternalID")
    file_index = _column(headers, "TrackFilepath")
    fields = ["Label", "AlbumNo", "AlbumTitle", "WorkTitle", "TrackNo", "workAudioId", "Filename"]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({
                "Label": "BMAT Custom",
                "AlbumNo": row[cat_index],
                "AlbumTitle": row[title_index],
                "WorkTitle": row[track_index],
                "TrackNo": row[number_index],
                "workAudioId": row[id_index],
                "Filename": PurePosixPath(str(row[file_index])).name,
            })


def prepare_delivery(
    submission_path: Path,
    releases_path: Path,
    root: Path,
    ledger_path: Path,
    logger: logging.Logger,
    *,
    delivery_date: date | None = None,
    include_catalogues: Iterable[str] = (),
    workflow_id: str | None = None,
) -> PreparedDelivery | None:
    ledger = load_ledger(ledger_path)
    selected = list(include_catalogues) or pending_catalogues(releases_path, ledger)
    for delivery in reversed(ledger.get("deliveries", [])):
        retry_catalogues = tuple(str(value) for value in delivery.get("catalogues", []))
        if (
            delivery.get("status") in {"prepared", "failed"}
            and (not workflow_id or delivery.get("workflow_id") == workflow_id)
            and retry_catalogues
            and set(retry_catalogues).issubset(selected)
        ):
            prepared = _prepared_from_ledger(root, delivery)
            logger.info(
                f"  Resuming BMAT batch {prepared.batch_id}: "
                f"{prepared.track_count} track(s), {len(prepared.catalogues)} catalogue(s)."
            )
            return prepared
    if not selected:
        logger.info(
            "  ✓ BMAT custom-release report has no new catalogues requiring delivery."
        )
        return None
    headers, rows = select_and_validate_submission(submission_path, selected)
    batch_id = _next_batch_id(root, ledger, delivery_date or date.today())
    package = root / batch_id
    package.mkdir(parents=True, exist_ok=False)
    metadata = package / f"{batch_id[:8]}_{int(batch_id[-4:]):02d}.xlsx"
    _write_metadata(metadata, headers, rows)
    request = root / "_WORKFLOW" / f"{batch_id}_dams_download_manifest.csv"
    request.parent.mkdir(parents=True, exist_ok=True)
    _write_dams_download_manifest(request, headers, rows)
    ledger["deliveries"].append({
        "batch_id": batch_id,
        "catalogues": selected,
        "track_count": len(rows),
        "status": "prepared",
        "attempts": [],
        "created_at": _utc_now(),
        "metadata_sha256": hashlib.sha256(metadata.read_bytes()).hexdigest(),
        "workflow_id": workflow_id,
    })
    save_ledger(ledger_path, ledger)
    logger.info(f"  Prepared BMAT batch {batch_id}: {len(rows)} track(s), {len(selected)} catalogue(s).")
    return PreparedDelivery(batch_id, package, metadata, tuple(selected), len(rows))


def _expected_files_by_catalogue(
    prepared: PreparedDelivery,
) -> dict[str, list[PurePosixPath]]:
    headers, rows = _load_xlsx_rows(prepared.metadata_path)
    catalogue_index = _column(headers, "CatNo", "AlbumNo")
    track_index = _column(headers, "TrackFilepath")
    artwork_index = _column(headers, "CDArtwork")
    expected = {catalogue: [] for catalogue in prepared.catalogues}
    for row_number, row in enumerate(rows, start=2):
        catalogue = str(row[catalogue_index] or "").strip()
        if catalogue not in expected:
            continue
        relative = _safe_remote_path(
            row[track_index], field="TrackFilepath", row_number=row_number
        )
        if relative is None or relative.parts[0] != catalogue:
            raise ValueError(
                f"Row {row_number}: TrackFilepath must begin with {catalogue}/"
            )
        expected[catalogue].append(relative)
        artwork = _safe_remote_path(
            row[artwork_index], field="CDArtwork", row_number=row_number
        )
        if artwork is not None:
            if artwork.parts[0] != catalogue:
                raise ValueError(
                    f"Row {row_number}: CDArtwork must begin with {catalogue}/"
                )
            expected[catalogue].append(artwork)
    for catalogue, paths in expected.items():
        expected[catalogue] = list(dict.fromkeys(paths))
    return expected


def _install_dams_archive(
    archive: Path,
    package: Path,
    expected: list[PurePosixPath],
) -> None:
    """Install exact manifest assets from a DAMS all-edits ZIP."""
    expected_names = {path.name for path in expected}
    if len(expected_names) != len(expected):
        raise ValueError("BMAT manifest contains duplicate asset basenames")
    with zipfile.ZipFile(archive) as bundle:
        file_members = [
            info for info in bundle.infolist()
            if not info.is_dir() and not info.filename.startswith("__MACOSX/")
        ]
        by_name: dict[str, list[zipfile.ZipInfo]] = {}
        for member in file_members:
            by_name.setdefault(PurePosixPath(member.filename).name, []).append(member)
        missing = sorted(expected_names - set(by_name))
        duplicates = sorted(name for name in expected_names if len(by_name.get(name, [])) != 1)
        unexpected_wavs = sorted(
            PurePosixPath(member.filename).name
            for member in file_members
            if PurePosixPath(member.filename).suffix.lower() == ".wav"
            and PurePosixPath(member.filename).name not in expected_names
        )
        if missing or duplicates or unexpected_wavs:
            raise ValueError(
                f"DAMS archive {archive.name} does not exactly match the BMAT "
                f"audio manifest; missing={missing}, duplicates={duplicates}, "
                f"unexpected_wavs={unexpected_wavs}"
            )
        for relative in expected:
            destination = package.joinpath(*relative.parts)
            if destination.is_file() and destination.stat().st_size > 0:
                continue
            destination.parent.mkdir(parents=True, exist_ok=True)
            temporary = destination.with_name(f".{destination.name}.dams-download.tmp")
            with bundle.open(by_name[relative.name][0]) as source, temporary.open("wb") as target:
                shutil.copyfileobj(source, target)
            if temporary.stat().st_size <= 44:
                temporary.unlink(missing_ok=True)
                raise ValueError(f"DAMS returned an empty/invalid WAV: {relative}")
            temporary.replace(destination)


def _dams_is_authenticated(page) -> bool:
    return (
        page.url.startswith(DAMS_BASE_URL)
        and "/albums" in page.url
        and page.locator('input[placeholder="Cat Number"]').count() == 1
    )


def _authenticate_dams(page, logger: logging.Logger, *, allow_interactive: bool) -> None:
    page.goto(DAMS_ALBUMS_URL, wait_until="domcontentloaded")
    if _dams_is_authenticated(page):
        logger.info("  DAMS authenticated using the private per-user session.")
        return

    # Normal runs may initiate the same UMG Employee SSO route once, but never
    # wait indefinitely for MFA. Interactive enrollment is confined to setup.
    employee_sso = page.get_by_text(
        re.compile(r"UMG.*Employee|Employee.*(Login|Sign.?in)", re.IGNORECASE)
    )
    if employee_sso.count() == 1:
        employee_sso.click()
    timeout = 10 * 60 * 1000 if allow_interactive else DAMS_AUTH_TIMEOUT_MS
    try:
        page.wait_for_function(
            """() => location.hostname === 'dams.universalproductionmusic.com'
                && location.pathname.startsWith('/albums')
                && document.querySelectorAll('input[placeholder="Cat Number"]').length === 1""",
            timeout=timeout,
            polling=500,
        )
    except Exception as exc:
        mode = "Complete UMG Employee SSO in the open browser" if allow_interactive else (
            "Run `python3 auth_manager.py --setup dams` outside the workflow"
        )
        raise RuntimeError(f"DAMS authentication did not complete. {mode}.") from exc
    logger.info("  DAMS authenticated using UMG Employee SSO.")


def setup_dams_auth(logger: logging.Logger) -> bool:
    """Enroll a private persistent DAMS browser session through UMG SSO."""
    from auth_manager import private_creation_umask, secure_private_directory
    try:
        from playwright.sync_api import sync_playwright
    except ModuleNotFoundError:
        logger.error(
            "DAMS setup requires Playwright. Install requirements.txt and the "
            "Chromium runtime first."
        )
        return False
    secure_private_directory(DAMS_PROFILE_DIR, recursive=True)
    try:
        with sync_playwright() as playwright:
            with private_creation_umask():
                context = playwright.chromium.launch_persistent_context(
                    user_data_dir=str(DAMS_PROFILE_DIR),
                    headless=False,
                    accept_downloads=False,
                )
            page = context.pages[0] if context.pages else context.new_page()
            try:
                _authenticate_dams(page, logger, allow_interactive=True)
                logger.info(
                    "Protected DAMS Albums page verified; leaving the browser "
                    "visible for 10 seconds for confirmation."
                )
                time.sleep(10)
            finally:
                context.close()
        secure_private_directory(DAMS_PROFILE_DIR, recursive=True)
        return True
    except Exception as exc:
        logger.error("DAMS authentication setup failed: %s", exc)
        return False


def _dams_album_audio_url(page, catalogue: str) -> str:
    # Every catalogue starts from the album list.  The preceding iteration
    # leaves the shared page on that album's Audio screen.
    page.goto(DAMS_ALBUMS_URL, wait_until="domcontentloaded")
    catalogue_filter = page.locator('input[placeholder="Cat Number"]')
    if catalogue_filter.count() != 1:
        raise RuntimeError("DAMS Albums page did not expose one Cat Number filter")
    catalogue_filter.fill(catalogue)
    page.wait_for_timeout(2_000)
    exact = re.compile(rf"^{re.escape(catalogue)}$")
    cells = page.get_by_text(exact)
    rows = cells.locator("xpath=ancestor::tr[1]")
    if rows.count() != 1:
        raise RuntimeError(
            f"DAMS catalogue {catalogue} resolved to {rows.count()} album rows; expected 1"
        )
    edit_links = rows.locator('a[href*="/albums/edit/"]')
    hrefs = [
        edit_links.nth(index).get_attribute("href")
        for index in range(edit_links.count())
    ]
    matches = [
        href for href in hrefs
        if href and re.search(r"/albums/edit/\d+$", href)
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"DAMS catalogue {catalogue} exposed {len(matches)} exact edit links; expected 1"
        )
    return f"{DAMS_BASE_URL}{matches[0]}/uploadAudio" if matches[0].startswith("/") else (
        re.sub(r"/+$", "", matches[0]) + "/uploadAudio"
    )


def _download_catalogue_archive(page, catalogue: str, destination: Path) -> Path:
    page.goto(_dams_album_audio_url(page, catalogue), wait_until="domcontentloaded")
    page.get_by_text(re.compile(rf"Catalogue number:\s*{re.escape(catalogue)}\b")).wait_for(
        timeout=DAMS_AUTH_TIMEOUT_MS
    )
    refresh = page.get_by_role("button", name=re.compile(r"^Refresh$"))
    if refresh.count() != 1:
        raise RuntimeError(f"DAMS Audio page for {catalogue} has no unique Refresh button")
    download_link = refresh.locator("xpath=following-sibling::a[1]")
    if download_link.count() != 1:
        raise RuntimeError(
            f"DAMS Audio page for {catalogue} has no unique album download control"
        )
    with page.expect_download(timeout=DAMS_DOWNLOAD_TIMEOUT_MS) as event:
        download_link.click()
    download = event.value
    suggested = Path(download.suggested_filename).name
    if not suggested.lower().endswith(".zip"):
        raise RuntimeError(f"DAMS download for {catalogue} was not a ZIP archive")
    destination.mkdir(parents=True, exist_ok=True)
    archive = destination / suggested
    download.save_as(str(archive))
    if not archive.is_file() or archive.stat().st_size == 0:
        raise RuntimeError(f"DAMS download for {catalogue} did not produce a file")
    return archive


def download_bmat_audio_from_dams(
    prepared: PreparedDelivery,
    root: Path,
    logger: logging.Logger,
    *,
    dry_run: bool,
) -> bool:
    expected = _expected_files_by_catalogue(prepared)
    missing_catalogues = [
        catalogue for catalogue, paths in expected.items()
        if any(not prepared.package_dir.joinpath(*path.parts).is_file() for path in paths)
    ]
    if not missing_catalogues:
        logger.info("  ✓ All BMAT DAMS audio is already present; download skipped.")
        return True
    if dry_run:
        logger.info(
            "  [DRY RUN] Would use UMG Employee SSO to open each DAMS Audio "
            f"page and download all edits for: {', '.join(missing_catalogues)}"
        )
        return True

    from auth_manager import private_creation_umask, secure_private_directory
    try:
        from playwright.sync_api import sync_playwright
    except ModuleNotFoundError:
        logger.error("  BMAT DAMS download requires Playwright and Chromium.")
        return False
    download_dir = root / "_WORKFLOW" / "downloads" / prepared.batch_id
    secure_private_directory(DAMS_PROFILE_DIR, recursive=True)
    try:
        with sync_playwright() as playwright:
            with private_creation_umask():
                context = playwright.chromium.launch_persistent_context(
                    user_data_dir=str(DAMS_PROFILE_DIR),
                    headless=False,
                    downloads_path=str(download_dir),
                    accept_downloads=True,
                )
            page = context.pages[0] if context.pages else context.new_page()
            try:
                _authenticate_dams(page, logger, allow_interactive=False)
                for catalogue in missing_catalogues:
                    logger.info(f"  ── DAMS custom album {catalogue} ──")
                    archive = _download_catalogue_archive(
                        page, catalogue, download_dir
                    )
                    _install_dams_archive(
                        archive, prepared.package_dir, expected[catalogue]
                    )
                    logger.info(
                        f"     ✓ Downloaded and installed {len(expected[catalogue])} manifest asset(s)."
                    )
            finally:
                context.close()
        secure_private_directory(DAMS_PROFILE_DIR, recursive=True)
        return True
    except Exception as exc:
        logger.error(f"  ✗ BMAT DAMS download failed: {exc}")
        return False


def verify_bmat_dams_audio(
    prepared: PreparedDelivery,
    logger: logging.Logger,
    *,
    dry_run: bool,
) -> bool:
    if dry_run:
        logger.info(
            f"  [DRY RUN] Would verify {prepared.track_count} WAV(s) downloaded "
            "from each custom album's DAMS Audio page."
        )
        return True
    try:
        normalize_package_layout(prepared.package_dir, prepared.metadata_path)
        validate_package(prepared.package_dir, prepared.metadata_path)
    except ValueError as exc:
        logger.error(
            "  BMAT audio is not ready. Download each custom album from its "
            f"DAMS Audio page, then retry: {exc}"
        )
        return False
    logger.info(
        f"  ✓ Verified {prepared.track_count} DAMS WAV(s) against the BMAT manifest."
    )
    return True


def normalize_package_layout(package: Path, metadata_path: Path) -> None:
    headers, rows = _load_xlsx_rows(metadata_path)
    track_index = _column(headers, "TrackFilepath")
    artwork_index = _column(headers, "CDArtwork")
    required = [
        path for row_number, row in enumerate(rows, start=2)
        for path in (
            _safe_remote_path(row[track_index], field="TrackFilepath", row_number=row_number),
            _safe_remote_path(row[artwork_index], field="CDArtwork", row_number=row_number),
        ) if path is not None
    ]
    for relative in dict.fromkeys(required):
        destination = package.joinpath(*relative.parts)
        if destination.is_file():
            continue
        matches = [p for p in package.rglob(relative.name) if p.is_file() and p != metadata_path]
        if len(matches) != 1:
            raise ValueError(f"Required package file has {len(matches)} source matches: {relative}")
        destination.parent.mkdir(parents=True, exist_ok=True)
        matches[0].replace(destination)


def validate_package(package: Path, metadata_path: Path) -> list[Path]:
    headers, rows = _load_xlsx_rows(metadata_path)
    track_index = _column(headers, "TrackFilepath")
    artwork_index = _column(headers, "CDArtwork")
    expected = {metadata_path.relative_to(package)}
    for row_number, row in enumerate(rows, start=2):
        for field, index in (("TrackFilepath", track_index), ("CDArtwork", artwork_index)):
            relative = _safe_remote_path(row[index], field=field, row_number=row_number)
            if relative is not None:
                expected.add(Path(*relative.parts))
    actual = {
        path.relative_to(package) for path in package.rglob("*")
        if path.is_file() and path.name not in {".DS_Store", "delivery.complete"}
    }
    missing = sorted(expected - actual)
    extra = sorted(actual - expected)
    if missing or extra:
        raise ValueError(f"BMAT manifest mismatch; missing={missing}, extra={extra}")
    empty = sorted(path for path in actual if (package / path).stat().st_size == 0)
    if empty:
        raise ValueError(f"BMAT package contains empty files: {empty}")
    return sorted(actual)


def _mkdir_remote(sftp, path: str) -> None:
    current = ""
    for part in PurePosixPath(path).parts:
        current = posixpath.join(current, part)
        try:
            sftp.stat(current)
        except OSError:
            sftp.mkdir(current)


def upload_package(
    prepared: PreparedDelivery,
    ledger_path: Path,
    logger: logging.Logger,
    *,
    known_hosts: Path | None = None,
) -> None:
    files = validate_package(prepared.package_dir, prepared.metadata_path)
    credentials = load_bmat_sftp_credentials()
    if credentials is None:
        raise RuntimeError("BMAT SFTP credentials are not available in this user's Keychain")
    username, password = credentials
    try:
        import paramiko
    except ImportError as exc:
        raise RuntimeError("BMAT upload requires Paramiko; install requirements.txt") from exc

    host_keys_path = known_hosts or (PRIVATE_STATE_DIR / "bmat_known_hosts")
    host_keys_path.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    client = paramiko.SSHClient()
    if host_keys_path.exists():
        client.load_host_keys(str(host_keys_path))
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    try:
        client.connect(
            BMAT_HOST, port=BMAT_PORT, username=username, password=password,
            look_for_keys=False, allow_agent=False, timeout=30,
        )
        # AutoAddPolicy accepts only a first-seen key. Persist it immediately;
        # a changed key on later runs is rejected by the default host-key check.
        client.save_host_keys(str(host_keys_path))
        os.chmod(host_keys_path, 0o600)
        sftp = client.open_sftp()
        remote_root = prepared.batch_id
        _mkdir_remote(sftp, remote_root)
        for relative in files:
            local = prepared.package_dir / relative
            remote = posixpath.join(remote_root, relative.as_posix())
            _mkdir_remote(sftp, posixpath.dirname(remote))
            temporary = remote + ".part"
            sftp.put(str(local), temporary)
            if sftp.stat(temporary).st_size != local.stat().st_size:
                raise RuntimeError(f"Remote size verification failed for {relative}")
            sftp.rename(temporary, remote)
        for relative in files:
            remote = posixpath.join(remote_root, relative.as_posix())
            if sftp.stat(remote).st_size != (prepared.package_dir / relative).stat().st_size:
                raise RuntimeError(f"Final remote verification failed for {relative}")
        with tempfile.NamedTemporaryFile() as marker:
            # This is intentionally the final remote mutation: BMAT begins
            # ingestion when the lower-case marker appears in the batch folder.
            sftp.put(marker.name, posixpath.join(remote_root, "delivery.complete"))
        sftp.close()
    finally:
        username = password = ""
        credentials = None
        client.close()

    ledger = load_ledger(ledger_path)
    for delivery in ledger["deliveries"]:
        if delivery.get("batch_id") == prepared.batch_id:
            delivery["status"] = "ingestion_pending"
            delivery.setdefault("attempts", []).append({"uploaded_at": _utc_now(), "result": "uploaded"})
            break
    save_ledger(ledger_path, ledger)
    logger.info(f"  ✓ Uploaded BMAT batch {prepared.batch_id}; delivery.complete was sent last.")


def _mark_delivery_failed(ledger_path: Path, batch_id: str, reason: str) -> None:
    """Record a retryable failure without storing sensitive exception text."""
    ledger = load_ledger(ledger_path)
    for delivery in ledger["deliveries"]:
        if delivery.get("batch_id") == batch_id:
            delivery["status"] = "failed"
            delivery.setdefault("attempts", []).append({
                "failed_at": _utc_now(),
                "result": reason,
            })
            break
    save_ledger(ledger_path, ledger)


def run_bmat_step(
    ctx: ReleaseContext,
    dry_run: bool,
    logger: logging.Logger,
    *,
    root: Path = BMAT_BASE,
) -> bool:
    """Run the date-scoped BMAT delivery as an independent workflow step."""
    from domo_exports import run_domo_exports

    logger.info(
        f"  BMAT release window: {ctx.release_start} → {ctx.release_end}\n"
        "  The local ledger excludes accepted and already-uploaded catalogues."
    )
    export_results = run_domo_exports(
        ctx,
        dry_run,
        logger,
        card_configs=BMAT_CARD_CONFIGS,
    )
    if dry_run:
        logger.info(
            "  [DRY RUN] Would select undelivered custom catalogues, prepare or "
            "resume one batch, download exact DAMS album audio, validate the "
            "package, upload it, and send delivery.complete last."
        )
        return True
    if not export_results or any(value != "ok" for value in export_results.values()):
        logger.error("  ✗ BMAT Domo exports did not both complete; delivery stopped.")
        return False

    ledger_path = root / "_WORKFLOW" / "delivery_ledger.json"
    prepared: PreparedDelivery | None = None
    try:
        prepared = prepare_delivery(
            ctx.bmat_submission_xlsx,
            ctx.bmat_releases_csv,
            root,
            ledger_path,
            logger,
            workflow_id=ctx.release_id,
        )
        if prepared is None:
            return True
        if not download_bmat_audio_from_dams(
            prepared, root, logger, dry_run=False
        ):
            _mark_delivery_failed(ledger_path, prepared.batch_id, "dams-download")
            return False
        if not verify_bmat_dams_audio(prepared, logger, dry_run=False):
            _mark_delivery_failed(ledger_path, prepared.batch_id, "package-validation")
            return False
        upload_package(prepared, ledger_path, logger)
        return True
    except Exception as exc:
        if prepared is not None:
            _mark_delivery_failed(ledger_path, prepared.batch_id, "step-error")
        logger.error(f"  ✗ BMAT delivery failed: {type(exc).__name__}: {exc}")
        return False


def seed_accepted_delivery(
    ledger_path: Path, batch_id: str, catalogues: Iterable[str], *, source: str
) -> None:
    ledger = load_ledger(ledger_path)
    if any(item.get("batch_id") == batch_id for item in ledger["deliveries"]):
        return
    ledger["deliveries"].append({
        "batch_id": batch_id,
        "catalogues": list(catalogues),
        "status": "accepted",
        "accepted_at": _utc_now(),
        "source": source,
        "attempts": [],
    })
    save_ledger(ledger_path, ledger)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Prepare a BMAT custom-content delivery")
    parser.add_argument("--submission", type=Path, required=True, help="BMAT submission XLSX export")
    parser.add_argument("--releases", type=Path, required=True, help="CUSTOM BMAT RELEASES CSV/XLSX export")
    parser.add_argument("--root", type=Path, default=BMAT_BASE)
    parser.add_argument("--include-catalog", action="append", default=[])
    parser.add_argument("--upload", action="store_true")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    logger = logging.getLogger("bmat")
    ledger_path = args.root / "_WORKFLOW" / "delivery_ledger.json"
    prepared = prepare_delivery(
        args.submission, args.releases, args.root, ledger_path, logger,
        include_catalogues=args.include_catalog,
    )
    if prepared is None:
        return 0
    if not verify_bmat_dams_audio(prepared, logger, dry_run=False):
        return 1
    if args.upload:
        upload_package(prepared, ledger_path, logger)
    else:
        logger.info("  Package validated locally; use --upload to send it to BMAT.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
