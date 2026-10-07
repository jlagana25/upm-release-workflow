"""Offline tests for BMAT package preparation and manifest validation."""

from __future__ import annotations

import tempfile
import unittest
import zipfile
import logging
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook

from bmat_delivery import (
    _install_dams_archive,
    _load_xlsx_rows,
    pending_catalogues,
    prepare_delivery,
    release_catalogues,
    validate_package,
)


class BmatDeliveryTests(unittest.TestCase):
    def setUp(self) -> None:
        self.logger = logging.getLogger("test-bmat-delivery")

    @staticmethod
    def _write_submission(path: Path, catalogue: str = "EMX1007") -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append([
            "CatNo", "InternalID", "TrackFilepath", "CDArtwork", "CDTitle",
            "TrackTitle", "TrackNo",
        ])
        worksheet.append([
            catalogue, "2260130", f"{catalogue}/track.wav", "", "Album",
            "Track", 1,
        ])
        workbook.save(path)

    def test_blank_trailing_artwork_cell_is_padded_and_valid(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            package = Path(tmp)
            metadata = package / "20260909_01.xlsx"
            headers = [
                "CatNo",
                "InternalID",
                "TrackFilepath",
                *[f"Unused {index}" for index in range(119)],
                "CDArtwork",
            ]
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(headers)
            worksheet.append([
                "EMX1007",
                "2260130",
                "EMX1007/track.wav",
            ])
            workbook.save(metadata)

            audio = package / "EMX1007" / "track.wav"
            audio.parent.mkdir()
            audio.write_bytes(b"RIFF-test")

            loaded_headers, rows = _load_xlsx_rows(metadata)
            self.assertEqual(len(loaded_headers), len(rows[0]))
            self.assertIsNone(rows[0][-1])
            self.assertEqual(
                [Path("20260909_01.xlsx"), Path("EMX1007/track.wav")],
                validate_package(package, metadata),
            )

    def test_pending_catalogues_excludes_accepted_and_uploaded_only(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            releases = Path(tmp) / "releases.csv"
            releases.write_text(
                "AlbumNo\nACCEPTED\nUPLOADED\nRETRY\nNEW\n", encoding="utf-8"
            )
            ledger = {
                "version": 1,
                "deliveries": [
                    {"status": "accepted", "catalogues": ["ACCEPTED"]},
                    {"status": "ingestion_pending", "catalogues": ["UPLOADED"]},
                    {"status": "failed", "catalogues": ["RETRY"]},
                ],
            }
            self.assertEqual(
                ["RETRY", "NEW"], pending_catalogues(releases, ledger)
            )

    def test_empty_release_exports_are_successful_noop(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for name, content in (
                ("empty.csv", ""),
                ("bom.csv", "\ufeff"),
                ("header.csv", "AlbumNo\n"),
                ("no-data.csv", "No data available\n"),
            ):
                path = root / name
                path.write_text(content, encoding="utf-8")
                self.assertEqual([], release_catalogues(path), name)

            workbook_path = root / "empty.xlsx"
            workbook = Workbook()
            workbook.active.delete_rows(1, workbook.active.max_row)
            workbook.save(workbook_path)
            self.assertEqual([], release_catalogues(workbook_path))

    def test_prepare_resumes_matching_failed_workflow_batch(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "BMAT"
            root.mkdir()
            submission = Path(tmp) / "submission.xlsx"
            releases = Path(tmp) / "releases.csv"
            ledger_path = root / "_WORKFLOW" / "delivery_ledger.json"
            self._write_submission(submission)
            releases.write_text("AlbumNo\nEMX1007\n", encoding="utf-8")

            first = prepare_delivery(
                submission,
                releases,
                root,
                ledger_path,
                self.logger,
                delivery_date=date(2026, 9, 11),
                workflow_id="UPM20260901",
                package_root=Path(tmp) / "release" / "3-FINAL PACKAGING" / "BMAT",
            )
            self.assertIsNotNone(first)
            self.assertEqual("BMAT", first.package_dir.parent.name)
            ledger = json.loads(ledger_path.read_text(encoding="utf-8"))
            ledger["deliveries"][0]["status"] = "failed"
            ledger_path.write_text(json.dumps(ledger), encoding="utf-8")
            releases.write_text(
                "AlbumNo\nEMX1007\nEMX1008\n", encoding="utf-8"
            )

            resumed = prepare_delivery(
                submission,
                releases,
                root,
                ledger_path,
                self.logger,
                delivery_date=date(2026, 9, 12),
                workflow_id="UPM20260901",
            )
            self.assertEqual(first.batch_id, resumed.batch_id)
            self.assertEqual(first.package_dir, resumed.package_dir)
            self.assertEqual(1, len(json.loads(
                ledger_path.read_text(encoding="utf-8")
            )["deliveries"]))

    def test_dams_archive_installs_only_exact_manifest_wavs(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            archive = base / "album.zip"
            package = base / "package"
            with zipfile.ZipFile(archive, "w") as bundle:
                bundle.writestr("download/track.wav", b"RIFF" + b"x" * 100)
            _install_dams_archive(
                archive, package, [Path("EMX1007/track.wav")]
            )
            self.assertEqual(
                b"RIFF" + b"x" * 100,
                (package / "EMX1007" / "track.wav").read_bytes(),
            )

    def test_dams_archive_rejects_unexpected_wav(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            archive = base / "album.zip"
            with zipfile.ZipFile(archive, "w") as bundle:
                bundle.writestr("track.wav", b"RIFF" + b"x" * 100)
                bundle.writestr("other.wav", b"RIFF" + b"y" * 100)
            with self.assertRaisesRegex(ValueError, "unexpected_wavs"):
                _install_dams_archive(
                    archive,
                    base / "package",
                    [Path("EMX1007/track.wav")],
                )


if __name__ == "__main__":
    unittest.main()
