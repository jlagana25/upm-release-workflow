#!/usr/bin/env python3
"""Submit the two SoundExchange registrant batches with strict validation."""

from __future__ import annotations

import argparse
import csv
import io
import json
import logging
import re
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Protocol

from config import PRIVATE_STATE_DIR, ReleaseContext, context_from_cli_args
from delivery_common import (
    DeliverySafetyError,
    checkpoint,
    collect_manifest,
    latest_workflow_gates,
    load_endpoint_state,
    private_json,
    receipt,
    require_live_authorization,
)
from delivery_state import set_partner_status


PORTAL_URL = "https://sxdirect.soundexchange.com/login/?next=%2Fcatalog%2Fsubmit%2F"


class SoundExchangeDeliveryError(DeliverySafetyError):
    pass


@dataclass(frozen=True)
class Registrant:
    key: str
    name: str
    registrant_id: str
    rights_owner: str
    filename_prefix: str


@dataclass(frozen=True)
class ValidationEntry:
    entry_number: str
    title: str
    isrc: str
    valid: bool
    messages: tuple[str, ...] = ()


REGISTRANTS = (
    Registrant(
        "mgb", "Universal Music - Mgb Na Llc", "2178141344",
        "Universal Mgb_Na_Llc", "ISRC Ingest Form - MGB NA LLC - Part ",
    ),
    Registrant(
        "ztunes", "Universal Music – Z Tunes, Llc", "2178141458",
        "Firstcom Music", "ISRC Ingest Form - Z TUNES LLC - Part ",
    ),
)


def _history_csv_text_isrcs(text: str) -> tuple[str, ...]:
    """Parse a portal history CSV already held in memory."""
    reader = csv.DictReader(io.StringIO(text.lstrip("\ufeff"), newline=""))
    if not reader.fieldnames:
        raise SoundExchangeDeliveryError("SoundExchange history CSV has no header")
    normalized = {
        re.sub(r"[^a-z0-9]", "", str(name).casefold()): name
        for name in reader.fieldnames
    }
    header = next(
        (
            original for key, original in normalized.items()
            if key == "isrc" or key.startswith("soundrecordingisrc")
        ),
        None,
    )
    if header is None:
        raise SoundExchangeDeliveryError("SoundExchange history CSV has no ISRC column")
    values = tuple(
        re.sub(r"[^A-Z0-9]", "", str(row.get(header) or "").upper())
        for row in reader
    )
    if not values or any(not value for value in values):
        raise SoundExchangeDeliveryError("SoundExchange history CSV contains a blank ISRC")
    return values


def _history_csv_isrcs(path: Path) -> tuple[str, ...]:
    """Read the portal's generated history CSV and return canonical ISRCs."""
    return _history_csv_text_isrcs(path.read_text(encoding="utf-8-sig"))


class SoundExchangeGateway(Protocol):
    def require_authenticated(self) -> None: ...
    def select_registrant(self, registrant: Registrant) -> None: ...
    def pending_count(self) -> int: ...
    def bulk_import(self, workbook: Path) -> None: ...
    def wait_for_validation(self, expected_total: int, timeout_seconds: float) -> tuple[ValidationEntry, ...]: ...
    def submit_recordings(self) -> None: ...
    def verify_upload_history(self, registrant: Registrant, expected_isrcs: tuple[str, ...]) -> str | None: ...
    def close(self) -> None: ...


class PlaywrightSoundExchangeGateway:
    """Visible retained-session adapter. Selectors fail closed if the UI drifts."""

    def __init__(self, *, interactive_login: bool = False) -> None:
        try:
            from playwright.sync_api import sync_playwright
        except ImportError as exc:
            raise SoundExchangeDeliveryError("SoundExchange delivery requires Playwright") from exc
        profile = PRIVATE_STATE_DIR / "soundexchange_profile"
        profile.mkdir(parents=True, exist_ok=True, mode=0o700)
        self._pw = sync_playwright().start()
        self._context = self._pw.chromium.launch_persistent_context(
            str(profile), channel="chrome", headless=False
        )
        self._page = self._context.pages[0] if self._context.pages else self._context.new_page()
        self._page.goto(PORTAL_URL, wait_until="domcontentloaded")
        self._interactive_login = interactive_login

    def require_authenticated(self) -> None:
        def ready() -> bool:
            return (
                not self._page.locator("input[type=password]:visible").count()
                and bool(self._page.get_by_text("MY CATALOG", exact=True).count())
            )

        if ready():
            return
        # Osano can be the only rendered UI immediately after navigation. If
        # login recovery runs during that gap, the generic form helper sees no
        # credential fields and can mistake the cookie controls for an
        # ambiguous sign-in action. Close only Osano's exact dialog control,
        # then give the real SoundExchange password field a bounded chance to
        # render before examining the form.
        cookie_close = self._page.get_by_text("Close this dialog", exact=True)
        if cookie_close.count() == 1 and cookie_close.is_visible():
            cookie_close.click()
        try:
            self._page.locator("input[type=password]:visible").first.wait_for(
                state="visible", timeout=15_000
            )
        except Exception:
            # `attempt_keychain_login` remains the authoritative fail-closed
            # form/ready-state check and emits the redacted setup guidance.
            pass
        from auth_manager import load_soundexchange_credentials
        from portal_auth import PortalAuthenticationError, attempt_keychain_login
        try:
            if attempt_keychain_login(
                self._page, load_soundexchange_credentials(), ready=ready
            ):
                return
        except PortalAuthenticationError as exc:
            raise SoundExchangeDeliveryError(
                f"SoundExchange authentication failed closed: {exc}"
            ) from exc
        if self._interactive_login:
            try:
                self._page.get_by_text("MY CATALOG", exact=True).wait_for(
                    state="visible", timeout=300_000
                )
            except Exception as exc:
                raise SoundExchangeDeliveryError(
                    "SoundExchange interactive sign-in did not complete within five minutes"
                ) from exc
            if ready():
                return
        raise SoundExchangeDeliveryError(
            "SoundExchange session is not authenticated; run "
            "auth_manager.py --enroll-soundexchange-keychain"
        )

    def select_registrant(self, registrant: Registrant) -> None:
        self._page.goto(
            "https://sxdirect.soundexchange.com/catalog/submit/",
            wait_until="domcontentloaded",
        )
        row = self._page.locator("tr", has_text=registrant.name)
        try:
            row.first.wait_for(state="visible", timeout=15_000)
        except Exception as exc:
            raise SoundExchangeDeliveryError(
                f"Could not load the verified registrant row for {registrant.name}"
            ) from exc
        if row.count() != 1 or registrant.registrant_id not in row.inner_text():
            raise SoundExchangeDeliveryError(f"Could not uniquely verify {registrant.name}")
        link = row.get_by_text(registrant.rights_owner, exact=True)
        if link.count() != 1:
            raise SoundExchangeDeliveryError(f"Rights Owner changed for {registrant.name}")
        link.click()
        self._page.locator(
            'input[type=file][data-cy="bulk-upload-file"]'
        ).wait_for(state="attached")

    def pending_count(self) -> int:
        payload = self._summary_records()
        return len(payload)

    def bulk_import(self, workbook: Path) -> None:
        picker = self._page.locator('input[type=file][data-cy="bulk-upload-file"]')
        if picker.count() != 1:
            raise SoundExchangeDeliveryError("Bulk Import file input was not unique")
        # Do not click the visible Bulk Import button first: it opens a native
        # file chooser that remains modal while Playwright separately injects
        # the file into the verified input. The portal imports immediately when
        # this exact input receives the workbook.
        picker.set_input_files(str(workbook))

    def _summary_records(self) -> list[dict[str, object]]:
        body = self._page.locator('[data-cy="summary-table-body"]')
        if body.count() != 1:
            return []
        return body.evaluate(
            """root => {
                const norm = value => String(value || '').toLowerCase().replace(/[^a-z0-9]/g, '');
                const controllers = [];
                let node = root;
                for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                    if (!window.angular) break;
                    const wrapped = window.angular.element(node);
                    for (const accessor of ['scope', 'isolateScope', 'data']) {
                        try {
                            const candidate = wrapped[accessor] && wrapped[accessor]();
                            if (candidate && candidate.ctrl) controllers.push(candidate.ctrl);
                            if (candidate && candidate.$ngControllerController) {
                                controllers.push(candidate.$ngControllerController);
                            }
                        } catch (_) {}
                    }
                }
                const seen = new WeakSet();
                const arrays = [];
                const walk = (value, depth) => {
                    if (!value || depth > 5 || (typeof value !== 'object')) return;
                    if (seen.has(value)) return;
                    seen.add(value);
                    if (Array.isArray(value)) {
                        if (value.some(item => item && typeof item === 'object' &&
                            Object.keys(item).some(key => norm(key).includes('isrc')))) arrays.push(value);
                        for (const item of value) walk(item, depth + 1);
                        return;
                    }
                    for (const [key, item] of Object.entries(value)) {
                        if (!key.startsWith('$') && key !== 'window' && key !== 'document') {
                            walk(item, depth + 1);
                        }
                    }
                };
                const direct = controllers.find(controller =>
                    controller && Array.isArray(controller.recordings)
                );
                const roots = [root, ...controllers];
                if (window.angular) {
                    const wrapped = window.angular.element(root);
                    try { roots.push(wrapped.scope()); } catch (_) {}
                    try { roots.push(wrapped.isolateScope()); } catch (_) {}
                    try { roots.push(wrapped.data()); } catch (_) {}
                }
                for (const candidate of roots) walk(candidate, 0);
                // The current portal exposes the authoritative ui-scroll
                // collection as ctrl.recordings. Do not select the largest
                // arbitrary Angular array: `countries` is larger and caused
                // completed imports to look permanently empty.
                const records = direct ? direct.recordings :
                    (arrays.sort((a, b) => b.length - a.length)[0] || []);
                const valueFor = (record, names) => {
                    for (const [key, value] of Object.entries(record || {})) {
                        if (names.includes(norm(key))) return value;
                    }
                    return undefined;
                };
                return records.map((record, index) => {
                    const isrc = valueFor(record, ['isrc', 'soundrecordingisrc']);
                    const title = valueFor(record, ['title', 'recordingtitle', 'soundrecordingtitle']);
                    const valid = valueFor(record, ['valid', 'isvalid', 'validationvalid']);
                    const processing = valueFor(record, ['processing', 'isprocessing', 'validating']);
                    const rawMessages = valueFor(record, ['messages', 'errors', 'validationmessages']);
                    const messages = Array.isArray(rawMessages) ? rawMessages.map(String) :
                        (rawMessages ? [String(rawMessages)] : []);
                    return {
                        entry_number: String(index + 1), title: String(title || ''),
                        isrc: String(isrc || ''), valid: valid === true,
                        ready: typeof valid === 'boolean' && processing !== true,
                        valid_type: typeof valid,
                        processing_type: typeof processing,
                        schema: Object.keys(record || {}).map(norm).filter(Boolean).sort(),
                        messages,
                    };
                });
            }"""
        )

    def wait_for_validation(self, expected_total: int, timeout_seconds: float) -> tuple[ValidationEntry, ...]:
        deadline = time.monotonic() + timeout_seconds
        while time.monotonic() < deadline:
            payload = self._summary_records()
            if len(payload) == expected_total and payload and all(
                item.get("valid_type") == "undefined" for item in payload
            ):
                schema = ", ".join(map(str, payload[0].get("schema") or ()))
                raise SoundExchangeDeliveryError(
                    "SoundExchange validation model changed; no boolean validity "
                    f"field was found. Record schema: {schema or 'empty'}"
                )
            if len(payload) == expected_total and all(item.get("ready") for item in payload):
                return tuple(
                    ValidationEntry(
                        str(item["entry_number"]), str(item["title"]),
                        re.sub(r"[^A-Z0-9]", "", str(item["isrc"]).upper()),
                        bool(item["valid"]), tuple(map(str, item.get("messages") or ())),
                    )
                    for item in payload
                )
            time.sleep(2)
        raise SoundExchangeDeliveryError("SoundExchange validation timed out")

    def submit_recordings(self) -> None:
        button = self._page.locator('[data-cy="submit-recordings-btn"]')
        if button.count() != 1 or not button.is_enabled():
            raise SoundExchangeDeliveryError("Submit Recordings is unavailable")
        button.click()

    def verify_upload_history(
        self, registrant: Registrant, expected_isrcs: tuple[str, ...]
    ) -> str | None:
        self._page.goto(
            "https://sxdirect.soundexchange.com/catalog/history/",
            wait_until="domcontentloaded",
        )
        owner_row = self._page.locator("tr", has_text=registrant.name)
        owner_row = owner_row.filter(has_text=registrant.registrant_id)
        try:
            owner_row.first.wait_for(state="visible", timeout=15_000)
        except Exception:
            return None
        if owner_row.count() != 1:
            return None
        action = owner_row.get_by_text("View Submission History", exact=True)
        if action.count() != 1:
            return None
        action.click()
        try:
            self._page.wait_for_url("**/catalog/history/*/", timeout=15_000)
        except Exception:
            return None
        expected = tuple(re.sub(r"[^A-Z0-9]", "", value.upper()) for value in expected_isrcs)
        rows = self._page.locator('tr[data-cy^="upload-row-"]')
        try:
            rows.first.wait_for(state="visible", timeout=15_000)
        except Exception:
            return None
        uploads = self._page.evaluate(
            """() => {
                let node = document.querySelector('tr[data-cy^="upload-row-"]');
                for (let depth = 0; node && depth < 8; depth += 1, node = node.parentElement) {
                    const wrapped = window.angular && window.angular.element(node);
                    for (const accessor of ['scope', 'isolateScope', 'data']) {
                        try {
                            const candidate = wrapped && wrapped[accessor] && wrapped[accessor]();
                            const controller = candidate && (
                                candidate.uploadHistoryPage || candidate.ctrl ||
                                candidate.$ngControllerController
                            );
                            if (controller && Array.isArray(controller.uploads)) {
                                return controller.uploads.slice(0, 25).map(upload => ({
                                    file_id: upload.fileId,
                                    filename: String(upload.userFileName || ''),
                                    date: String(upload.date || ''),
                                    type: String(upload.type || ''),
                                    isrc_download: upload.isrcDownload === true,
                                }));
                            }
                        } catch (_) {}
                    }
                }
                return [];
            }"""
        )
        for upload in uploads:
            if not upload.get("isrc_download") or upload.get("file_id") is None:
                continue
            csv_text = self._page.evaluate(
                """async fileId => {
                    const token = String(window.csrfmiddlewaretoken || '');
                    const body = new URLSearchParams({
                        fileId: String(fileId),
                        'X-CSRFToken': token,
                        csrfmiddlewaretoken: token,
                    });
                    const response = await fetch('/catalog/api/isrc/csv/', {
                        method: 'POST',
                        credentials: 'same-origin',
                        headers: {
                            'Content-Type': 'application/x-www-form-urlencoded;charset=UTF-8',
                            'X-CSRFToken': token,
                        },
                        body: body.toString(),
                    });
                    if (!response.ok) throw new Error(`history csv http ${response.status}`);
                    return await response.text();
                }""",
                upload["file_id"],
            )
            actual = _history_csv_text_isrcs(str(csv_text))
            if len(actual) != len(expected) or sorted(actual) != sorted(expected):
                continue
            return (
                f"{registrant.key}:{upload['file_id']}:{len(actual)}:"
                f"{upload.get('date', '')}:{str(upload.get('filename', ''))[:80]}"
            )
        return None

    def close(self) -> None:
        self._context.close()
        self._pw.stop()


def _part_number(path: Path) -> int:
    match = re.search(r" - Part (\d+)\.xlsx$", path.name, re.I)
    if not match:
        raise SoundExchangeDeliveryError(f"Unexpected SoundExchange workbook name: {path.name}")
    return int(match.group(1))


def registrant_workbooks(ctx: ReleaseContext, registrant: Registrant) -> tuple[Path, ...]:
    root = ctx.soundexchange_final_dir
    paths = tuple(sorted(root.glob(registrant.filename_prefix + "*.xlsx"), key=_part_number))
    if not paths:
        raise SoundExchangeDeliveryError(f"No final workbooks for {registrant.name}")
    parts = [_part_number(path) for path in paths]
    if parts != list(range(1, len(parts) + 1)):
        raise SoundExchangeDeliveryError(f"Non-contiguous workbook parts for {registrant.name}: {parts}")
    return paths


def workbook_rows(path: Path) -> tuple[tuple[str, str], ...]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SoundExchangeDeliveryError("SoundExchange validation requires openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        isrc_index = None
        title_index = None
        header_row = None
        # SoundExchange's official ingest template has nine title/instruction
        # rows before its real row-10 headers. Generated workbooks preserve
        # that layout, so locate the header by content instead of assuming row 1.
        for row_number, row in enumerate(rows, start=1):
            normalized = {
                re.sub(r"[^a-z0-9]", "", str(value or "").casefold()): index
                for index, value in enumerate(row)
                if str(value or "").strip()
            }
            isrc_index = next(
                (
                    index for key, index in normalized.items()
                    if key == "isrc"
                    or (
                        key.startswith("isrc")
                        and "ingest" not in key
                        and "file" not in key
                    )
                    or key.startswith("soundrecordingisrc")
                ),
                None,
            )
            if isrc_index is None:
                if row_number >= 50:
                    break
                continue
            title_index = next(
                (
                    index for key, index in normalized.items()
                    if key == "title" or key.startswith("tracktitle")
                    or key.startswith("recordingtitle")
                    or key.startswith("soundrecordingtitle")
                ),
                None,
            )
            header_row = row_number
            break
        if isrc_index is None or header_row is None:
            raise SoundExchangeDeliveryError(
                f"Workbook has no ISRC header in its first 50 rows: {path}"
            )
        result: list[tuple[str, str]] = []
        for row in rows:
            if not any(value not in (None, "") for value in row):
                continue
            isrc = str(row[isrc_index] or "").strip().upper()
            title = str(row[title_index] or "").strip() if title_index is not None else ""
            if not isrc:
                raise SoundExchangeDeliveryError(f"Workbook contains a blank ISRC: {path}")
            result.append((isrc, title))
        return tuple(result)
    finally:
        workbook.close()


def _invalid_audit(
    ctx: ReleaseContext,
    registrant: Registrant,
    workbooks: tuple[Path, ...],
    entries: tuple[ValidationEntry, ...],
) -> Path:
    invalid = [entry for entry in entries if not entry.valid]
    workbook_by_isrc = {
        isrc: workbook.name
        for workbook in workbooks
        for isrc, _title in workbook_rows(workbook)
    }
    path = ctx.specials_dir / "_WORKFLOW" / "soundexchange_invalid_entries" / f"{registrant.key}.json"
    return private_json(path, {
        "registrant": registrant.name,
        "registrant_id": registrant.registrant_id,
        "rights_owner": registrant.rights_owner,
        "workbooks": [item.name for item in workbooks],
        "invalid_count": len(invalid),
        "total_count": len(entries),
        "created_at": datetime.now(timezone.utc).isoformat(),
        "entries": [
            entry.__dict__ | {"source_workbook": workbook_by_isrc.get(entry.isrc)}
            for entry in invalid
        ],
    })


def deliver_soundexchange(
    ctx: ReleaseContext,
    dry_run: bool,
    logger: logging.Logger,
    *,
    gateway: SoundExchangeGateway | None = None,
    live_confirmation: str | None = None,
    validation_timeout_seconds: float = 3600,
    interactive_login: bool = False,
) -> bool:
    owned = False
    try:
        batches = [(registrant, registrant_workbooks(ctx, registrant)) for registrant in REGISTRANTS]
        all_files = tuple(
            item for _, paths in batches for path in paths
            for item in collect_manifest(path, allowed_suffixes=frozenset({".xlsx"}))
        )
        expected = {
            registrant.key: tuple(row for path in paths for row in workbook_rows(path))
            for registrant, paths in batches
        }
        for registrant, _ in batches:
            isrcs = [isrc for isrc, _ in expected[registrant.key]]
            if len(isrcs) != len(set(isrcs)):
                raise SoundExchangeDeliveryError(f"Duplicate ISRC for {registrant.name}")
        all_isrcs = [isrc for rows in expected.values() for isrc, _ in rows]
        if len(all_isrcs) != len(set(all_isrcs)):
            raise SoundExchangeDeliveryError("An ISRC appears in both registrant batches")
        if dry_run:
            logger.info("  [DRY RUN] SoundExchange: %d workbook(s), %d recording(s)", len(all_files), sum(map(len, expected.values())))
            return True
        require_live_authorization(ctx, live_confirmation)
        gate_ok, detail = latest_workflow_gates(ctx, ("10 SoundExchange forms",))
        if not gate_ok:
            raise SoundExchangeDeliveryError(f"SoundExchange submission is blocked: {detail}")
        if gateway is None:
            gateway = PlaywrightSoundExchangeGateway(interactive_login=interactive_login)
            owned = True
        gateway.require_authenticated()
        history_ids: dict[str, str] = {}
        for registrant, paths in batches:
            rows = expected[registrant.key]
            expected_isrcs = tuple(isrc for isrc, _title in rows)
            existing_history = gateway.verify_upload_history(registrant, expected_isrcs)
            if existing_history:
                history_ids[registrant.key] = existing_history
                continue
            prior = load_endpoint_state(ctx, "soundexchange")
            uncertain_submit = any(
                item.get("phase") == "submit_intent" and item.get("detail") == registrant.name
                for item in ((prior or {}).get("history") or [])
            )
            if uncertain_submit:
                raise SoundExchangeDeliveryError(
                    f"{registrant.name} has an unverified prior Submit Recordings action; "
                    "refusing a duplicate"
                )
            gateway.select_registrant(registrant)
            if gateway.pending_count() != 0:
                raise SoundExchangeDeliveryError(
                    f"{registrant.name} has pending recordings from another run"
                )
            for path in paths:
                gateway.bulk_import(path)
            entries = gateway.wait_for_validation(len(rows), validation_timeout_seconds)
            actual_isrcs = [entry.isrc for entry in entries]
            expected_isrcs_list = [isrc for isrc, _ in rows]
            invalid = [entry for entry in entries if not entry.valid]
            if invalid:
                audit = _invalid_audit(ctx, registrant, paths, entries)
                raise SoundExchangeDeliveryError(
                    f"{len(invalid)} invalid entries for {registrant.name}; audit: {audit}"
                )
            if len(entries) != len(rows) or sorted(actual_isrcs) != sorted(expected_isrcs_list):
                raise SoundExchangeDeliveryError(
                    f"Validated recordings do not exactly match {registrant.name} workbooks"
                )
            checkpoint(ctx, "soundexchange", all_files, "validated", registrant.name)
            checkpoint(
                ctx, "soundexchange", all_files, "submit_intent", registrant.name,
                remote_ids=history_ids.values(),
            )
            gateway.submit_recordings()
            deadline = time.monotonic() + validation_timeout_seconds
            history_id = gateway.verify_upload_history(registrant, expected_isrcs)
            while not history_id and time.monotonic() < deadline:
                time.sleep(2)
                history_id = gateway.verify_upload_history(registrant, expected_isrcs)
            if not history_id:
                raise SoundExchangeDeliveryError(f"Upload History did not verify {registrant.name}")
            history_ids[registrant.key] = history_id
            checkpoint(
                ctx, "soundexchange", all_files, "submitted", registrant.name,
                remote_ids=history_ids.values(),
            )
        path = receipt(ctx, "soundexchange", all_files, {
            "registrants": [
                {
                    "name": registrant.name,
                    "registrant_id": registrant.registrant_id,
                    "rights_owner": registrant.rights_owner,
                    "workbooks": [item.name for item in paths],
                    "recordings": len(expected[registrant.key]),
                    "history_id": history_ids[registrant.key],
                }
                for registrant, paths in batches
            ]
        })
        set_partner_status(ctx.specials_dir, "soundexchange", "delivered")
        logger.info("  ✓ SoundExchange submissions verified: %s", path)
        return True
    except Exception as exc:
        logger.error("  ✗ SoundExchange delivery failed: %s", exc)
        return False
    finally:
        if owned and gateway is not None:
            gateway.close()


def _main() -> int:
    parser = argparse.ArgumentParser(description="Submit SoundExchange registrant workbooks")
    parser.add_argument("--year", type=int)
    parser.add_argument("--month", type=int)
    parser.add_argument("--part", type=int, choices=(1, 2))
    parser.add_argument("--previous-month", action="store_true")
    parser.add_argument("--start-date")
    parser.add_argument("--end-date")
    parser.add_argument("--full-month-content", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--confirm-live-release")
    parser.add_argument(
        "--interactive-login", action="store_true",
        help="wait up to five minutes for a one-session SoundExchange sign-in",
    )
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    return 0 if deliver_soundexchange(
        context_from_cli_args(args), args.dry_run, logging.getLogger("soundexchange"),
        live_confirmation=args.confirm_live_release,
        interactive_login=args.interactive_login,
    ) else 1


if __name__ == "__main__":
    raise SystemExit(_main())
